from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from backend.app.excel import build_workbook
from backend.app.models import ExportRequest, Lead


def test_excel_export_contains_leads_and_search_summary():
    lead = Lead(
        name="Plomberie Boréale", address="Québec", place_id="abc", zone_index=1,
        zone_latitude=46.8, zone_longitude=-71.2, collected_at=datetime.now(timezone.utc),
    )

    content = build_workbook(ExportRequest(leads=[lead], search={"query": "plombier"}))
    workbook = load_workbook(BytesIO(content))

    assert workbook.sheetnames == ["Leads", "Recherche"]
    assert workbook["Leads"]["A2"].value == "Plomberie Boréale"
    assert workbook["Leads"]["I2"].value == "abc"
    assert workbook["Recherche"]["B2"].value == "plombier"


def test_excel_export_neutralizes_formula_like_text():
    lead = Lead(
        name="=HYPERLINK(\"https://example.invalid\",\"Cliquez\")",
        address="@adresse",
        place_id="formula-test",
        zone_index=1,
        zone_latitude=46.8,
        zone_longitude=-71.2,
        collected_at=datetime.now(timezone.utc),
    )

    content = build_workbook(ExportRequest(leads=[lead], search={"query": "=1+1"}))
    workbook = load_workbook(BytesIO(content), data_only=False)

    assert workbook["Leads"]["A2"].data_type == "s"
    assert workbook["Leads"]["A2"].value.startswith("'=")
    assert workbook["Leads"]["B2"].data_type == "s"
    assert workbook["Leads"]["B2"].value == "'@adresse"
    assert workbook["Recherche"]["B2"].data_type == "s"
    assert workbook["Recherche"]["B2"].value == "'=1+1"
