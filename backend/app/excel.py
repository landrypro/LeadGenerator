from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ExportRequest

HEADERS = [
    "Name", "Address", "Phone", "InternationalPhone", "Website", "GoogleMapsUrl",
    "Latitude", "Longitude", "PlaceId", "PrimaryType", "BusinessStatus",
    "ServiceAreaBusiness", "ZoneIndex", "ZoneLatitude", "ZoneLongitude",
    "DistanceKm", "RadiusVerified", "CollectedAt",
]

DANGEROUS_EXCEL_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


def safe_excel_value(value: object) -> object:
    """Keep external text from being interpreted as an Excel formula."""
    if isinstance(value, str) and value.startswith(DANGEROUS_EXCEL_PREFIXES):
        return f"'{value}"
    return value


def build_workbook(payload: ExportRequest) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"
    sheet.append(HEADERS)
    for lead in payload.leads:
        row = [
            lead.name, lead.address, lead.phone, lead.international_phone, lead.website,
            lead.google_maps_url, lead.latitude, lead.longitude, lead.place_id,
            lead.primary_type, lead.business_status, lead.service_area_business,
            lead.zone_index, lead.zone_latitude, lead.zone_longitude, lead.distance_km,
            lead.radius_verified, lead.collected_at.isoformat(),
        ]
        sheet.append([safe_excel_value(value) for value in row])

    header_fill = PatternFill("solid", fgColor="163E34")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24
    widths = [28, 42, 18, 22, 34, 36, 12, 12, 30, 22, 18, 20, 12, 15, 15, 13, 15, 25]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in range(2, sheet.max_row + 1):
        for column in (5, 6):
            cell = sheet.cell(row, column)
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    summary = workbook.create_sheet("Recherche")
    summary.append(["Paramètre", "Valeur"])
    for key, value in payload.search.items():
        summary.append([safe_excel_value(key), safe_excel_value(str(value))])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 48
    summary.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
