#!/usr/bin/env python3
"""Build the Marketteo CRM user manual DOCX from its Markdown source."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


GREEN = RGBColor(0x17, 0x3F, 0x35)
GREEN_2 = RGBColor(0x1F, 0x5A, 0x49)
GOLD = RGBColor(0xE3, 0xA8, 0x29)
MINT = "D9EEE3"
CREAM = "F8F2E7"
INK = RGBColor(0x17, 0x25, 0x1F)
MUTED = RGBColor(0x6F, 0x7D, 0x76)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS.items():
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def ensure_child(parent, tag):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def apply_table_geometry(table, widths):
    widths = [int(width) for width in widths]
    if sum(widths) != TABLE_WIDTH_DXA:
        widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = ensure_child(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_ind = ensure_child(tbl_pr, "w:tblInd")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for column, width in zip(table.columns, widths):
        column.width = Twips(width)
    for row in table.rows:
        row.height = None
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_w = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcW")
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    text_run.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph._p.extend((begin, instr, separate, text_run, end))


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, GREEN_2, 18, 10),
        "Heading 2": (13, GREEN_2, 14, 7),
        "Heading 3": (12, GREEN, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_numbering(document):
    """Align built-in level-0 list definitions to the compact guide preset."""
    numbering = document.part.numbering_part.element
    for abstract in numbering.findall(qn("w:abstractNum")):
        level = abstract.find(qn("w:lvl"))
        if level is None or level.get(qn("w:ilvl")) != "0":
            continue
        paragraph_style = level.find(qn("w:pStyle"))
        if paragraph_style is None or paragraph_style.get(qn("w:val")) not in {"ListBullet", "ListNumber"}:
            continue
        p_pr = ensure_child(level, "w:pPr")
        tabs = ensure_child(p_pr, "w:tabs")
        tab = tabs.find(qn("w:tab"))
        if tab is None:
            tab = OxmlElement("w:tab")
            tabs.append(tab)
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        indent = ensure_child(p_pr, "w:ind")
        indent.set(qn("w:left"), "540")
        indent.set(qn("w:hanging"), "270")


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True


def add_running_furniture(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("MARKETTEO CRM  |  MANUEL UTILISATEUR")
    set_font(r, size=8.5, color=MUTED, bold=True)
    p_pr = p._p.get_or_add_pPr()
    borders = ensure_child(p_pr, "w:pBdr")
    bottom = ensure_child(borders, "w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), "D9EEE3")

    footer = section.footer
    page_p = footer.paragraphs[0]
    page_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    page_p.paragraph_format.space_before = Pt(4)
    run = page_p.add_run("Version 0.5  |  Page ")
    set_font(run, size=8.5, color=MUTED)
    add_field(page_p, "PAGE")


def add_cover(document):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(112)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GUIDE DE RÉFÉRENCE")
    set_font(r, size=10, color=GOLD, bold=True)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Marketteo CRM")
    set_font(r, size=31, color=GREEN, bold=True)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("Manuel utilisateur")
    set_font(r, size=17, color=GREEN_2)

    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(34)
    rule_pr = rule._p.get_or_add_pPr()
    borders = ensure_child(rule_pr, "w:pBdr")
    bottom = ensure_child(borders, "w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "E3A829")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("Version 0.5 - Édition illustrée à valider")
    set_font(r, size=11, color=GREEN, bold=True)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("État fonctionnel du 26 août 2026")
    set_font(r, size=10, color=MUTED, italic=True)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(88)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Commerciaux  |  Gestionnaires  |  Administrateurs")
    set_font(r, size=9.5, color=MUTED)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_toc(document, headings):
    title = document.add_paragraph("Sommaire", style="Heading 1")
    title.paragraph_format.space_before = Pt(0)
    for heading in headings:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(heading)
        set_font(r, size=10.5, color=GREEN)
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    r = note.add_run("Conseil : utilisez le volet Navigation de Word pour parcourir les titres du manuel.")
    set_font(r, size=9.5, color=MUTED, italic=True)
    note.add_run().add_break(WD_BREAK.PAGE)


INLINE_PATTERN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
IMAGE_PATTERN = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
CAPTION_PATTERN = re.compile(r"^\*(Figure\s+\d+\s+—.+)\*$")


def add_inline(paragraph, text):
    for part in INLINE_PATTERN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Consolas", size=9.5, color=GREEN)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), CREAM)
            run._element.get_or_add_rPr().append(shading)
        else:
            run = paragraph.add_run(part)
            set_font(run)


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(cells)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def table_widths(rows):
    columns = len(rows[0])
    if columns == 5 and rows[0][0] == "Fonction":
        return [3000, 1100, 1800, 1600, 1860]
    if columns == 3:
        return [2600, 3380, 3380]
    if columns == 4 and rows[0][0] == "Version":
        return [1000, 1600, 1700, 5060]
    weights = [max(8, max(len(row[i]) if i < len(row) else 0 for row in rows)) for i in range(columns)]
    total = sum(weights)
    widths = [round(TABLE_WIDTH_DXA * weight / total) for weight in weights]
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    return widths


def add_markdown_table(document, rows):
    if not rows:
        return
    columns = len(rows[0])
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(paragraph, values[column_index] if column_index < len(values) else "")
            for run in paragraph.runs:
                set_font(run, size=9.3, bold=True if row_index == 0 else run.bold)
            if row_index == 0:
                set_cell_shading(cell, MINT)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F8FAF8")
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, table_widths(rows))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(document, text):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shading = ensure_child(p_pr, "w:shd")
    shading.set(qn("w:fill"), CREAM)
    borders = ensure_child(p_pr, "w:pBdr")
    left = ensure_child(borders, "w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "E3A829")
    add_inline(p, text)
    for run in p.runs:
        set_font(run, size=10, color=GREEN)


def add_figure(document, image_path, alt_text, caption):
    """Add a centered, accessible figure without letting tall images overflow a page."""
    if not image_path.is_file():
        raise FileNotFoundError(f"Image introuvable : {image_path}")
    from PIL import Image

    with Image.open(image_path) as image:
        pixels_width, pixels_height = image.size
    ratio = pixels_height / pixels_width
    height_inches = min(6.25 * ratio, 5.6)
    width_inches = height_inches / ratio

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(width_inches), height=Inches(height_inches))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", alt_text)

    caption_p = document.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(10)
    caption_p.paragraph_format.keep_together = True
    caption_run = caption_p.add_run(caption)
    set_font(caption_run, size=9, color=MUTED, italic=True)


def add_body_from_markdown(document, markdown, source_dir):
    lines = markdown.splitlines()
    start = next((i for i, line in enumerate(lines) if line.startswith("## 1.")), 0)
    index = start
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        image_match = IMAGE_PATTERN.match(line)
        if image_match:
            caption = ""
            if index + 1 < len(lines):
                caption_match = CAPTION_PATTERN.match(lines[index + 1].strip())
                if caption_match:
                    caption = caption_match.group(1)
                    index += 1
            add_figure(document, source_dir / image_match.group(2), image_match.group(1), caption)
            index += 1
            continue
        if line.startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(document, rows)
            continue
        if line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("## "):
            document.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("> "):
            add_callout(document, line[2:])
        elif re.match(r"^\d+\. ", line):
            p = document.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            p = document.add_paragraph()
            add_inline(p, line)
        index += 1


def build(source, output):
    markdown = source.read_text(encoding="utf-8")
    headings = [line[3:].strip() for line in markdown.splitlines() if line.startswith("## ")]
    transition_note = next(
        (line[2:].strip() for line in markdown.splitlines() if line.startswith("> ")),
        "",
    )
    document = Document()
    configure_styles(document)
    configure_numbering(document)
    configure_section(document.sections[0])
    add_running_furniture(document.sections[0])
    document.core_properties.title = "Marketteo CRM - Manuel utilisateur"
    document.core_properties.subject = "Guide d'utilisation de Marketteo CRM"
    document.core_properties.author = "Marketteo"
    document.core_properties.keywords = "Marketteo, CRM, manuel utilisateur"
    add_cover(document)
    add_toc(document, headings)
    if transition_note:
        add_callout(document, transition_note)
    add_body_from_markdown(document, markdown, source.parent)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()
